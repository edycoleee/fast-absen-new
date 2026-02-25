from fastapi import APIRouter, Depends, Query, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from io import BytesIO
from config.database import get_db
from schemas.user import UserCreate, UserUpdate
from services.user_service import UserService
from repositories.user_repository import UserRepository
from utils.response import success_response, paginated_response
from utils.dependencies import CommonQueryParams, require_permission
from utils.permission_registry import PermissionKeys
from utils.constants import ErrorMessages, HTTPStatus as HC, SuccessMessages
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/users", tags=["Users"])

_TEMPLATE_HEADERS = [
    ("username",   "Username *",       "tst_john"),
    ("password",   "Password",          "Absen@1234"),
    ("id_pegawai", "ID Pegawai",        "P001"),
    ("role_names", "Role (pisahkan koma)", "admin,pegawai"),
    ("is_active",  "Is Active (TRUE/FALSE)", "TRUE"),
]


def _build_template_workbook() -> BytesIO:
    """Build the Excel import template and return it as a BytesIO stream."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    example_fill = PatternFill("solid", fgColor="D9E1F2")

    # Header row
    for col, (_, label, _) in enumerate(_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Example row
    for col, (_, _, example) in enumerate(_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=example)
        cell.fill = example_fill

    # Column widths
    widths = [20, 18, 16, 30, 22]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ─── NOTE ON ROUTE ORDER ─────────────────────────────────────────────────────
# template/download and import must be declared BEFORE /{user_id} so FastAPI
# does not treat "template" or "import" as integer path parameters.
# ─────────────────────────────────────────────────────────────────────────────


# ─── GET /users/template/download ────────────────────────────────────────────
@router.get("/template/download", summary="Download template Excel untuk import users")
async def download_template(
    _=Depends(require_permission(PermissionKeys.USERS_CREATE)),
):
    buf = _build_template_workbook()
    headers = {"Content-Disposition": 'attachment; filename="users_import_template.xlsx"'}
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


# ─── POST /users/import ───────────────────────────────────────────────────────
@router.post("/import", summary="Import users dari file Excel (.xlsx)")
async def import_users(
    file: UploadFile = File(..., description="File Excel (.xlsx) — gunakan template download"),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_CREATE)),
):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(HC.BAD_REQUEST, "File harus berformat .xlsx")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(HC.BAD_REQUEST, "File Excel tidak valid atau rusak")

    ws = wb.active
    raw_headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]

    # map display header → field name
    header_map = {label: field for field, label, _ in _TEMPLATE_HEADERS}
    col_to_field = {i: header_map.get(h, h.split(" ")[0].lower()) for i, h in enumerate(raw_headers)}

    rows = []
    for ws_row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in ws_row):
            continue
        row_dict = {col_to_field[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(ws_row)}
        rows.append(row_dict)

    wb.close()
    if not rows:
        raise HTTPException(HC.BAD_REQUEST, "File tidak memiliki data (hanya header)")

    svc = UserService(UserRepository(db))
    result = await svc.import_users(rows)
    return success_response(
        f"Import selesai: {result['created']} berhasil, "
        f"{result['skipped']} dilewati, {len(result['errors'])} gagal",
        data=result,
    )


# ─── GET /users/ ──────────────────────────────────────────────────────────────
@router.get("")
async def list_users(
    q: CommonQueryParams = Depends(),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_READ)),
):
    svc = UserService(UserRepository(db))
    items, total = await svc.get_all(page=q.page, limit=q.limit)
    data = [{
        "id": u.id, "username": u.username,
        "id_pegawai": u.id_pegawai, "is_active": u.is_active,
    } for u in items]
    return paginated_response("Berhasil", items=data, page=q.page, limit=q.limit, total=total)


# ─── POST /users/ ─────────────────────────────────────────────────────────────
@router.post("", status_code=HC.CREATED)
async def create_user(
    body: UserCreate,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_CREATE)),
):
    svc = UserService(UserRepository(db))
    user = await svc.create_user(
        {"id_pegawai": body.id_pegawai, "username": body.username, "password": body.password},
        body.role_ids or [],
    )
    return success_response(SuccessMessages.CREATED.format("User"), data={"id": user.id, "username": user.username})


# ─── GET /users/{user_id} ─────────────────────────────────────────────────────
@router.get("/{user_id}")
async def get_user(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_READ)),
):
    svc = UserService(UserRepository(db))
    user = await svc.get_by_id_with_roles(user_id)
    if not user:
        raise HTTPException(status_code=HC.NOT_FOUND, detail=ErrorMessages.NOT_FOUND.format("User"))
    return success_response("Berhasil", data={
        "id": user.id, "username": user.username,
        "id_pegawai": user.id_pegawai, "is_active": user.is_active,
        "roles": [r.name for r in user.roles],
        "permissions": list(user.permissions),
    })


# ─── PUT /users/{user_id} ─────────────────────────────────────────────────────
@router.put("/{user_id}")
async def update_user(
    user_id: int,
    body: UserUpdate,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_UPDATE)),
):
    svc = UserService(UserRepository(db))
    user = await svc.update_user(user_id, body.model_dump(exclude_none=True), body.role_ids)
    return success_response(SuccessMessages.UPDATED.format("User"), data={"id": user.id, "username": user.username})


# ─── DELETE /users/{user_id} ──────────────────────────────────────────────────
@router.delete("/{user_id}")
async def delete_user(
    user_id: int,
    force: bool = Query(False, description="Hapus permanen (true) atau nonaktifkan (false)"),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.USERS_DELETE)),
):
    svc = UserService(UserRepository(db))
    await svc.delete_user(user_id, force=force)
    msg = SuccessMessages.DELETED.format("User") if force else "User berhasil dinonaktifkan."
    return success_response(msg)

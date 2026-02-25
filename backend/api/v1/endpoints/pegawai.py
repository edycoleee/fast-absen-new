from fastapi import APIRouter, Depends, Query, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from io import BytesIO
from typing import Optional
from config.database import get_db
from schemas.pegawai import PegawaiCreate, PegawaiUpdate
from services.pegawai_service import PegawaiService
from repositories.pegawai_repository import PegawaiRepository
from utils.response import success_response, paginated_response
from utils.dependencies import CommonQueryParams, require_permission
from utils.permission_registry import PermissionKeys
from utils.constants import ErrorMessages, HTTPStatus as HC, SuccessMessages
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/pegawai", tags=["Pegawai"])

_TEMPLATE_HEADERS = [
    ("id_pegawai",    "ID Pegawai *",              "P001"),
    ("nama",          "Nama *",                     "Budi Santoso"),
    ("nip",           "NIP",                        "198501012010011001"),
    ("jenis_kelamin", "Jenis Kelamin (L/P)",        "L"),
    ("tempat_lahir",  "Tempat Lahir",               "Jakarta"),
    ("tanggal_lahir", "Tanggal Lahir (YYYY-MM-DD)", "1985-01-01"),
    ("alamat",        "Alamat",                     "Jl. Merdeka No. 1"),
    ("id_unit",       "ID Unit",                    "1"),
    ("kepala_id_unit","Kepala ID Unit",              ""),
    ("is_active",     "Is Active (TRUE/FALSE)",      "TRUE"),
]


def _build_template_workbook() -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pegawai"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    example_fill = PatternFill("solid", fgColor="D9E1F2")

    for col, (_, label, _) in enumerate(_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for col, (_, _, example) in enumerate(_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=example)
        cell.fill = example_fill

    widths = [14, 24, 22, 20, 18, 24, 30, 10, 14, 22]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _fmt(p):
    return {
        "id_pegawai": p.id_pegawai, "nip": p.nip, "nama": p.nama,
        "id_unit": p.id_unit, "kepala_id_unit": p.kepala_id_unit,
        "jenis_kelamin": p.jenis_kelamin,
        "tempat_lahir": p.tempat_lahir, "tanggal_lahir": str(p.tanggal_lahir) if p.tanggal_lahir else None,
        "alamat": p.alamat, "is_active": p.is_active, "foto": p.foto,
    }


# ─── GET /pegawai/template/download ───────────────────────────────────────────
@router.get("/template/download", summary="Download template Excel untuk import pegawai")
async def download_template(
    _=Depends(require_permission(PermissionKeys.PEGAWAI_CREATE)),
):
    buf = _build_template_workbook()
    headers = {"Content-Disposition": 'attachment; filename="pegawai_import_template.xlsx"'}
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


# ─── POST /pegawai/import ───────────────────────────────────────────────────────
@router.post("/import", summary="Import pegawai dari file Excel (.xlsx)")
async def import_pegawai(
    file: UploadFile = File(..., description="File Excel (.xlsx) — gunakan template download"),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_CREATE)),
):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(HC.BAD_REQUEST, "File harus berformat .xlsx")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(HC.BAD_REQUEST, "File Excel tidak valid atau rusak")

    ws = wb.active
    raw_headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]

    # map display label → field name
    label_to_field = {label: field for field, label, _ in _TEMPLATE_HEADERS}
    col_to_field = {
        i: label_to_field.get(h, h.split(" ")[0].lower().rstrip("*").strip())
        for i, h in enumerate(raw_headers)
    }

    rows = []
    for ws_row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in ws_row):
            continue
        row_dict = {
            col_to_field[i]: (str(v).strip() if v is not None else "")
            for i, v in enumerate(ws_row)
            if i in col_to_field
        }
        rows.append(row_dict)

    wb.close()
    if not rows:
        raise HTTPException(HC.BAD_REQUEST, "File tidak memiliki data (hanya header)")

    svc = PegawaiService(PegawaiRepository(db))
    result = await svc.import_pegawai(rows)
    return success_response(
        f"Import selesai: {result['created']} berhasil, "
        f"{result['skipped']} dilewati, {len(result['errors'])} gagal",
        data=result,
    )


# ─── GET /pegawai/ ─────────────────────────────────────────────────────────────
@router.get("")
async def list_pegawai(
    q: CommonQueryParams = Depends(),
    unit_id: Optional[int] = Query(None),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_READ)),
):
    svc = PegawaiService(PegawaiRepository(db))
    items, total = await svc.get_paged(q.page, q.limit, q.search, unit_id)
    return paginated_response("Berhasil", items=[_fmt(p) for p in items], page=q.page, limit=q.limit, total=total)


# ─── POST /pegawai/ ────────────────────────────────────────────────────────────
@router.post("", status_code=HC.CREATED)
async def create_pegawai(
    body: PegawaiCreate,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_CREATE)),
):
    svc = PegawaiService(PegawaiRepository(db))
    p = await svc.create_pegawai(body.model_dump())
    return success_response(SuccessMessages.CREATED.format("Pegawai"), data=_fmt(p))


# ─── GET /pegawai/{id_pegawai} ──────────────────────────────────────────────────
@router.get("/{id_pegawai}")
async def get_pegawai(
    id_pegawai: str,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_READ)),
):
    svc = PegawaiService(PegawaiRepository(db))
    p = await svc.get_by_id(id_pegawai)
    if not p:
        raise HTTPException(HC.NOT_FOUND, ErrorMessages.NOT_FOUND.format("Pegawai"))
    return success_response("Berhasil", data=_fmt(p))


# ─── PUT /pegawai/{id_pegawai} ─────────────────────────────────────────────────
@router.put("/{id_pegawai}")
async def update_pegawai(
    id_pegawai: str,
    body: PegawaiUpdate,
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_UPDATE)),
):
    svc = PegawaiService(PegawaiRepository(db))
    p = await svc.update_pegawai(id_pegawai, body.model_dump())
    return success_response(SuccessMessages.UPDATED.format("Pegawai"), data=_fmt(p))


# ─── DELETE /pegawai/{id_pegawai} ───────────────────────────────────────────────
@router.delete("/{id_pegawai}")
async def delete_pegawai(
    id_pegawai: str,
    force: bool = Query(False, description="Hapus permanen (true) atau nonaktifkan (false)"),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_permission(PermissionKeys.PEGAWAI_DELETE)),
):
    svc = PegawaiService(PegawaiRepository(db))
    await svc.delete_pegawai(id_pegawai, force=force)
    msg = SuccessMessages.DELETED.format("Pegawai") if force else "Pegawai berhasil dinonaktifkan."
    return success_response(msg)
